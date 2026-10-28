"""
IMPORT TAP GMV -> tab "all" di "[TAP] TT <bulan>"
================================================================================
Satu file per kali (xlsx, tab "Custom report"). Cuma baris yang Creator-attributed GMV-nya
bukan 0 yang dibawa (laporan TAP itu cross-join, hampir semua barisnya 0).
Brand + Sub Brand dicari dari List Produk lewat Product ID (sama dengan import_aff.py).
Produk dengan Sub Brand di SKIP_SUBS (Lip Care) tidak dimasukkan.
Kolom B:AC ditulis ke bawah data yang ada; kolom A (penanda manual "batas 18") tidak disentuh.
TANPA CEK DUPLIKAT: jangan upload rentang tanggal yang sudah ada (app memberi peringatan).
Kolom M, Z:AC ditulis sebagai rumus, sama dengan baris yang sudah ada di sheet.
Ikut mengisi sheet gabungan AFF - TAP (import_comb.py) kecuali pakai --tanpa-gabungan.

TERMINAL:
    python import_tap.py "TAP GMV 19-25 OCT 2026.xlsx"            <- simulasi, tidak menulis apa-apa
    python import_tap.py "TAP GMV 19-25 OCT 2026.xlsx" --apply    <- benar-benar menulis
Product ID yang belum punya brand: dibuatkan Isi_Brand_<nama file>.xlsx, isi lalu jalankan lagi.
"""
import argparse
import collections
import datetime
import re
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import openpyxl

import import_aff as ia
import import_comb as ic

TARGET_ID, TARGET_TAB = '1HfwlqHLKzZ11f9LCxwhqzUfm_kGeHtqzardBLFMFrak', 'all'
SOURCE_TAB = 'Custom report'
DATE_RE = re.compile(r'\d{4}-\d{2}-\d{2}$')
GMV = 'Creator-attributed GMV'
# kolom file mentah yang dipakai
NEED = ['Date', 'Campaign name', 'Creator name', 'Creator follower count', 'Product ID', 'Product name', GMV,
        'Affiliate video-attributed GMV', 'Creator LIVE-attributed GMV', 'Revenue (Showcase)', 'Link GMV',
        'Creator-attributed orders', 'Creator LIVE-attributed orders', 'Creator video-attributed orders',
        'LIVE likes', 'Video likes', 'Video views', 'LIVE views', 'LIVE streams', 'Videos',
        'Creator-attributed items sold']


NS = '{http://schemas.openxmlformats.org/spreadsheetml/2006/main}'
RID = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id'


def read_raw(path):
    """-> (baris dengan GMV != 0 sebagai dict header -> nilai, jumlah baris yang dipindai).
    xlsx dibaca langsung dari XML-nya: jauh lebih hemat RAM dan cepat daripada openpyxl untuk file 90 MB /
    500 ribu baris. Kalau susunan file tidak terduga, jatuh ke openpyxl."""
    try:
        with zipfile.ZipFile(path) as z:
            return _collect(_stream_rows(z))
    except zipfile.BadZipFile:
        raise ValueError('File ini bukan xlsx yang valid.')
    except (KeyError, ET.ParseError):
        wb = openpyxl.load_workbook(path, read_only=True)
        try:
            ws = wb[SOURCE_TAB] if SOURCE_TAB in wb.sheetnames else wb.worksheets[0]
            return _collect(ws.iter_rows(values_only=True))
        finally:
            wb.close()


def _cell(el, strings):
    t, v = el.get('t'), el.find(NS + 'v')
    if t == 'inlineStr':
        return ''.join(x.text or '' for x in el.iter(NS + 't'))
    if v is None or v.text is None:
        return None
    if t == 's':
        return strings[int(v.text)]
    if t in ('str', 'e'):
        return v.text
    if t == 'b':
        return v.text == '1'
    f = float(v.text)
    return int(f) if f.is_integer() else f


def _to_list(row):
    """{'AB12': nilai, ...} -> list nilai, posisi = nomor kolom."""
    out = []
    for ref, v in row.items():
        n = 0
        for ch in ref:
            if not ch.isalpha():
                break
            n = n * 26 + ord(ch) - 64
        out.extend([None] * (n - len(out)))
        out[n - 1] = v
    return out


def _stream_rows(z):
    """Baris sheet "Custom report" (atau sheet pertama) sebagai list nilai, dibaca bertahap tanpa memuat semuanya."""
    strings = []
    if 'xl/sharedStrings.xml' in z.namelist():
        with z.open('xl/sharedStrings.xml') as f:
            for _, el in ET.iterparse(f):
                if el.tag == NS + 'si':
                    strings.append(''.join(t.text or '' for t in el.iter(NS + 't')))
                    el.clear()
    sheets = list(ET.fromstring(z.read('xl/workbook.xml')).find(NS + 'sheets'))
    rels = {r.get('Id'): r.get('Target') for r in ET.fromstring(z.read('xl/_rels/workbook.xml.rels'))}
    sheet = next((s for s in sheets if s.get('name') == SOURCE_TAB), sheets[0])
    target = rels[sheet.get(RID)]
    row = {}
    with z.open(target[1:] if target.startswith('/') else 'xl/' + target) as f:
        for _, el in ET.iterparse(f):
            if el.tag == NS + 'c':
                row[el.get('r')] = _cell(el, strings)
                el.clear()
            elif el.tag == NS + 'row':
                yield _to_list(row)
                row = {}
                el.clear()


def _collect(rows_iter):
    rows_iter = iter(rows_iter)
    header = [str(h).strip() if h is not None else '' for h in next(rows_iter, [])]
    missing = [c for c in NEED if c not in header]
    if missing:
        raise ValueError(f'Kolom ini tidak ada di file: {", ".join(missing)}. Bukan laporan TAP GMV?')
    idx = [header.index(c) for c in NEED]
    rows, scanned = [], 0
    width = max(idx) + 1
    for r in rows_iter:
        r = list(r) + [None] * (width - len(r))
        if r[idx[0]] is None:
            continue
        scanned += 1
        row = dict(zip(NEED, (r[i] for i in idx)))
        try:
            if ia.num(row[GMV]) in ('', 0):
                continue
        except ValueError:
            raise ValueError(f'Baris {scanned + 1}: {GMV} bukan angka ({row[GMV]!r}).')
        if isinstance(row['Date'], (int, float)) and 30000 < row['Date'] < 80000:  # tanggal Excel (serial)
            row['Date'] = datetime.date(1899, 12, 30) + datetime.timedelta(days=int(row['Date']))
        if isinstance(row['Date'], (datetime.date, datetime.datetime)):
            row['Date'] = row['Date'].strftime('%Y-%m-%d')
        row['Date'] = str(row['Date']).strip()
        if not DATE_RE.match(row['Date']):
            raise ValueError(f'Baris {scanned + 1}: format Date tidak dikenal ({row["Date"]!r}), harusnya YYYY-MM-DD.')
        row['Product ID'] = str(row['Product ID']).strip()
        rows.append(row)
    return rows, scanned


def build_row(r, n, brand, sub):
    """28 nilai kolom B..AC untuk baris sheet nomor n."""
    name, campaign = (' '.join(str(r[k]).split()) for k in ('Product name', 'Campaign name'))
    return [
        r['Date'], campaign, ia.txt(r['Creator name']), ia.num(r['Creator follower count']),
        ia.txt(r['Product ID']), ia.txt(name), brand, sub,
        ia.num(r[GMV]), ia.num(r['Affiliate video-attributed GMV']), ia.num(r['Creator LIVE-attributed GMV']),
        f'=sum(N{n}:O{n})', ia.num(r['Revenue (Showcase)']), ia.num(r['Link GMV']),
        ia.num(r['Creator-attributed orders']), ia.num(r['Creator LIVE-attributed orders']),
        ia.num(r['Creator video-attributed orders']),
        ia.num(r['LIVE likes']), ia.num(r['Video likes']), ia.num(r['Video views']), ia.num(r['LIVE views']),
        ia.num(r['LIVE streams']), ia.num(r['Videos']), ia.num(r['Creator-attributed items sold']),
        f'=Y{n}/P{n}', f'=R{n}*Z{n}', f'=Q{n}*Z{n}', f'=Y{n}-sum(AA{n}:AB{n})',
    ]


def read_dates(svc, sid=TARGET_ID):
    """(baris terakhir terisi di kolom B, Counter tanggal ISO -> jumlah baris) dari tab all."""
    vals = svc.values().get(spreadsheetId=sid, range=f"'{TARGET_TAB}'!B2:B",
                            valueRenderOption='UNFORMATTED_VALUE').execute(num_retries=3).get('values', [])
    base = datetime.date(1899, 12, 30)
    dates = collections.Counter(
        (base + datetime.timedelta(days=int(v[0]))).isoformat()
        for v in vals if v and isinstance(v[0], (int, float)))
    return 1 + len(vals), dates


def plan_rows(svc, rows, mapping, sid=TARGET_ID):
    """-> (baris pertama, baris terakhir sekarang, nilai B..AC yang akan ditulis, Counter tanggal yang sudah ada)."""
    last_row, dates = read_dates(svc, sid)
    first = last_row + 1
    out = [build_row(r, first + j, *mapping[r['Product ID']]) for j, r in enumerate(rows)]
    return first, last_row, out, dates


def write_rows(svc, out, first, sid=TARGET_ID):
    if not out:
        return 0
    ia.ensure_rows(svc, sid, TARGET_TAB, first + len(out) - 1)
    for s in range(0, len(out), 2000):
        chunk = out[s:s + 2000]
        svc.values().update(
            spreadsheetId=sid, range=f"'{TARGET_TAB}'!B{first + s}:AC{first + s + len(chunk) - 1}",
            valueInputOption='USER_ENTERED', body={'values': chunk}).execute(num_retries=3)
        print(f'Tertulis {s + len(chunk)}/{len(out)} baris.')
    return len(out)


def main(argv=None):
    ap = argparse.ArgumentParser(description='Import laporan TAP GMV (xlsx) ke tab "all" dan sheet gabungan.')
    ap.add_argument('file')
    ap.add_argument('--apply', action='store_true', help='benar-benar menulis ke Google Sheets')
    ap.add_argument('--mapping', help='file isian Brand/Sub Brand (default Isi_Brand_<nama file>.xlsx)')
    ap.add_argument('--tanpa-gabungan', action='store_true', help='jangan ikut update sheet gabungan AFF - TAP')
    ap.add_argument('--hanya-gabungan', action='store_true',
                    help='cuma update sheet gabungan (sheet utama sudah terisi, misal tadi gabungan gagal)')
    a = ap.parse_args(argv)

    src = Path(a.file)
    print(f'Membaca {src.name} (file besar, bisa 2-3 menit)...')
    try:
        rows, scanned = read_raw(src)
    except ValueError as e:
        sys.exit(f'ERROR: {e}')
    if not rows:
        sys.exit('ERROR: tidak ada baris dengan GMV selain 0 di file ini.')
    print(f'{scanned} baris dipindai, {len(rows)} baris dengan GMV selain 0.')
    names, count = {}, collections.Counter()
    for r in rows:
        names.setdefault(r['Product ID'], r['Product name'])
        count[r['Product ID']] += 1
    mapping_file = Path(a.mapping) if a.mapping else src.with_name(f'Isi_Brand_{src.stem}.xlsx')

    svc = ia.get_service(a.apply)
    mapping, new_ids, unknown_in_sheet = ia.resolve(svc, list(names), names, count, mapping_file)
    for p, row in unknown_in_sheet.items():
        if p in names:
            print(f'CATATAN: {p} ada di List Produk baris {row} dengan brand "{ia.UNKNOWN}"; '
                  'isian dipakai untuk import ini, tapi baris itu tidak diubah script (ubah manual).')

    rows, skipped = ia.drop_skipped(rows, mapping)
    if skipped:
        print(f'Dilewati (Sub Brand {sorted(ia.SKIP_SUBS)}): {len(skipped)} baris dari '
              f'{len({r["Product ID"] for r in skipped})} produk, tidak dimasukkan.')
    print(f'Baca tab {TARGET_TAB} untuk cari baris terakhir...')
    first, last_row, out, dates = plan_rows(svc, rows, mapping)

    in_file = collections.Counter(r['Date'] for r in rows)
    gmv = sum(float(ia.num(r[GMV]) or 0) for r in rows)
    by_brand = collections.Counter(mapping[r['Product ID']][0] for r in rows)
    print(f'\nBaris: {len(out)} -> baris {first}-{first + len(out) - 1} di {TARGET_TAB} (baris terakhir sekarang {last_row})')
    print(f'Tanggal: {", ".join(f"{d} ({n})" for d, n in sorted(in_file.items()))}')
    print(f'All GMV: {gmv:,.0f}'.replace(',', '.'))
    print('Per brand:', dict(by_brand))
    overlap = {d: dates[d] for d in in_file if dates.get(d)}
    if overlap:
        print(f'PERINGATAN: tab {TARGET_TAB} sudah berisi {sum(overlap.values())} baris untuk {len(overlap)} tanggal yang '
              f'juga ada di file ini ({min(overlap)} s/d {max(overlap)}). Kalau dilanjut, datanya dobel.')
    print(f'ID baru untuk List Produk: {len(new_ids)}')
    for p, name, brand, sub in new_ids:
        print(f'  {p} | {name[:60]} | {brand} | {sub}')

    comb_base = None if a.tanpa_gabungan else ic.tap_rows(rows, mapping)
    if not a.apply:
        if comb_base is not None:
            ic.run(svc, comb_base, 'TAP', False)
        print('\nDRY-RUN: tidak ada yang ditulis. Tambahkan --apply untuk menulis.')
        return
    if not out and not new_ids:
        print('\nTidak ada yang perlu ditulis.')
        return

    if not a.hanya_gabungan:
        ia.write_new_ids(svc, new_ids)
        write_rows(svc, out, first)
        ia.CACHE_FILE.write_text(ia.json.dumps(mapping, ensure_ascii=False), encoding='utf-8')
        print('Sheet utama selesai.')
    if comb_base is not None:
        try:
            ic.run(svc, comb_base, 'TAP', True)
        except ia.HttpError as e:
            sys.exit(f'ERROR: sheet utama sudah selesai, tapi gagal update sheet gabungan: {ic.explain(e)}\n'
                     'Setelah diperbaiki, jalankan lagi dengan --hanya-gabungan.')
    print('\nSELESAI.')


if __name__ == '__main__':
    main()
