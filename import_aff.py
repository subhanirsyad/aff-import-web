"""
IMPORT AFFILIATE TIKTOK (AFF GAL / PONDS / UHBS / VSL ...) -> tab ALL di "[AFF] TT <bulan>"
================================================================================
Satu file per run (CSV atau XLSX), semua brand sama-sama 33 kolom jadi tidak dibedakan.
TANPA FILTER: semua baris dibawa (semua Order Status, semua Platform).

ALUR
1. python import_aff.py "AFF PONDS 19-25 OCT 2026.csv"          <- dry-run, TIDAK menulis apa-apa
2. Kalau ada Product ID yang belum punya brand, script berhenti dan membuat
   "Isi_Brand_<nama file>.xlsx" (Product Name sudah terisi). Isi kolom Brand + Sub Brand,
   lalu jalankan perintah yang sama lagi (file itu otomatis terbaca).
3. Kalau hasil dry-run sudah benar: tambahkan --apply
   python import_aff.py "AFF PONDS 19-25 OCT 2026.csv" --apply
   --apply menambah ID baru ke tab "List Produk" lalu baris data ke bawah tab ALL.
   Tambah --limit 20 untuk cuma menulis 20 baris pertama (buat cek tampilan dulu).

Brand dicari dari brand_map.json (cache) -> kalau tidak ada, baca ulang "List Produk"
-> kalau masih tidak ada, minta diisi manual. Hapus brand_map.json kalau mau paksa baca ulang.
Tanpa cek duplikat: semua baris file ditambahkan ke bawah ALL. Jangan upload file yang sama dua kali.
Kolom Category (D) dan kolom hitung (B, M, S-Z) ditulis sebagai rumus.
Produk dengan Sub Brand di SKIP_SUBS (sekarang: Lip Care) tidak dimasukkan ke ALL; jumlahnya dicatat di log.
"""
import argparse
import collections
import csv
import datetime
import json
import re
import sys
from pathlib import Path

import openpyxl
from google.oauth2 import service_account
from googleapiclient.discovery import build
from openpyxl.worksheet.datavalidation import DataValidation

HERE = Path(__file__).resolve().parent
CREDENTIALS_FILE = HERE / 'credentials.json'
CACHE_FILE = HERE / 'brand_map.json'

TARGET_ID, TARGET_TAB = '1M9CdQYr_fTInUg7WewfPz6081qAHQyrBgsRqyCdmOm0', 'ALL'
BRAND_ID, BRAND_TAB = '1Wi0kcm7apS4gbLDcHJdz8P1iRohF62aGzM38Oy_f7gM', 'List Produk'
UNKNOWN = 'Tidak Diketahui'
SKIP_SUBS = {'lip care'}  # Sub Brand yang tidak dimasukkan ke ALL (huruf kecil), baris-barisnya dilewati

CONTENT_TYPE = {
    'Video': '2. Video', 'Livestream': '1. Live Streaming',
    'Showcase': '3. Sharelink', 'External Traffic Program': '3. Sharelink',
}
NEED = ['Product ID', 'Product Name', 'Price', 'Payment Amount', 'Quantity', 'Order Status',
        'Creator Username', 'Content Type', 'Content ID', 'Standard commission rate',
        'Est. Commission Base', 'Time Created']
TIME_RE = re.compile(r'^\d{2}/\d{2}/\d{4} \d{2}:\d{2}:\d{2}$')


# ---------------------------------------------------------------- baca file
def cell_str(v):
    if v is None:
        return ''
    if isinstance(v, datetime.datetime):
        return v.strftime('%d/%m/%Y %H:%M:%S')
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip().lstrip("'")


def read_raw(path):
    path = Path(path)
    ext = path.suffix.lower()
    if ext == '.csv':
        with open(path, encoding='utf-8-sig', newline='') as f:
            rows = list(csv.DictReader(f))
    elif ext in ('.xlsx', '.xlsm'):
        it = openpyxl.load_workbook(path, read_only=True, data_only=True).worksheets[0].iter_rows(values_only=True)
        head = [cell_str(c) for c in next(it)]
        rows = [dict(zip(head, map(cell_str, r))) for r in it if any(c not in (None, '') for c in r)]
    else:
        sys.exit(f'ERROR: format {ext} tidak didukung (pakai .csv atau .xlsx).')
    if not rows:
        sys.exit('ERROR: file kosong.')
    missing = [c for c in NEED if c not in rows[0]]
    if missing:
        sys.exit(f'ERROR: kolom tidak ketemu di header: {", ".join(missing)}')
    bad_ct = sorted({r['Content Type'] for r in rows} - set(CONTENT_TYPE))
    if bad_ct:
        sys.exit(f'ERROR: Content Type baru {bad_ct}. Tambahkan ke CONTENT_TYPE di import_aff.py '
                 '(1. Live Streaming / 2. Video / 3. Sharelink).')
    bad_t = next((r['Time Created'] for r in rows if not TIME_RE.match(r['Time Created'])), None)
    if bad_t is not None:
        sys.exit(f'ERROR: format Time Created tidak dikenal: "{bad_t}" (harus dd/mm/yyyy hh:mm:ss).')
    return rows


# ---------------------------------------------------------------- nilai sel
def num(s):
    s = str(s).strip().replace(',', '')
    if s == '':
        return ''
    f = float(s)
    return int(f) if f.is_integer() else f


def txt(s):
    """Apostrof di depan supaya Sheets simpan sebagai teks (ID 16+ digit, nama diawali = + - @)."""
    s = str(s)
    if s and (s[0] in '=+-@' or (s.isdigit() and len(s) >= 16) or TIME_RE.match(s)):
        return "'" + s
    return s


def week_of(time_created):
    d = int(time_created[:2])
    return 'Week 1' if d <= 7 else 'Week 2' if d <= 14 else 'Week 3' if d <= 21 else 'Week 4' if d <= 28 else 'Week 5'


def build_row(r, n, brand, sub):
    """26 nilai kolom A..Z untuk baris sheet nomor n."""
    creator_f = [f"=IFERROR(VLOOKUP(O{n},'LIST krEATOR'!${c}:${c},1,0))" for c in 'ABCDEFG']
    return [
        txt(r['Time Created']), f'=LEFT(A{n},10)', week_of(r['Time Created']),
        f'=VLOOKUP(G{n},Sheet4!B:C,2,0)',
        txt(r['Product ID']), txt(r['Product Name']), brand, sub,
        num(r['Price']), num(r['Payment Amount']), num(r['Est. Commission Base']), num(r['Quantity']),
        f'=L{n}*I{n}', r['Order Status'], txt(r['Creator Username']), CONTENT_TYPE[r['Content Type']],
        txt(r['Content ID']), r['Standard commission rate'],
        '=' + '+'.join(f'COUNTA({c}{n})' for c in 'TUVWXYZ'),
        *creator_f,
    ]


def drop_skipped(rows, mapping):
    """Buang baris yang Sub Brand-nya ada di SKIP_SUBS -> (baris yang dipakai, baris yang dibuang)."""
    skip = lambda r: mapping[r['Product ID']][1].strip().lower() in SKIP_SUBS
    return [r for r in rows if not skip(r)], [r for r in rows if skip(r)]


# ---------------------------------------------------------------- Google Sheets
def get_service(write, info=None):
    """info = isi service account sebagai dict (versi web); kosong -> credentials.json."""
    scope = 'https://www.googleapis.com/auth/spreadsheets' if write else \
        'https://www.googleapis.com/auth/spreadsheets.readonly'
    if info:
        creds = service_account.Credentials.from_service_account_info(dict(info), scopes=[scope])
    else:
        creds = service_account.Credentials.from_service_account_file(str(CREDENTIALS_FILE), scopes=[scope])
    return build('sheets', 'v4', credentials=creds).spreadsheets()


def grid_rows(svc, sid, tab):
    meta = svc.get(spreadsheetId=sid, fields='sheets(properties(sheetId,title,gridProperties))').execute()
    for s in meta['sheets']:
        if s['properties']['title'] == tab:
            return s['properties']['sheetId'], s['properties']['gridProperties']['rowCount']
    sys.exit(f'ERROR: tab "{tab}" tidak ketemu.')


def ensure_rows(svc, sid, tab, last_row_needed):
    tab_id, have = grid_rows(svc, sid, tab)
    if last_row_needed > have:
        svc.batchUpdate(spreadsheetId=sid, body={'requests': [{'appendDimension': {
            'sheetId': tab_id, 'dimension': 'ROWS', 'length': last_row_needed - have}}]}).execute()
        print(f'Grid tab {tab} ditambah {last_row_needed - have} baris.')


def read_brand_sheet(svc):
    vals = svc.values().get(spreadsheetId=BRAND_ID, range=f"'{BRAND_TAB}'!A2:G").execute().get('values', [])
    rows = []
    for i, r in enumerate(vals):
        r = [str(x).strip() for x in r] + [''] * (7 - len(r))
        if r[1]:
            rows.append(dict(row=i + 2, no=r[0], pid=r[1], name=r[2], brand=r[3], sub=r[4]))
    return rows, len(vals) + 1  # baris terakhir yang terisi


def sheet_id_from(link):
    """ID spreadsheet dari link Google Sheets (atau ID polos)."""
    m = re.search(r'/d/([\w-]+)', link) or re.fullmatch(r'\s*([\w-]{25,})\s*', link)
    if not m:
        raise ValueError('Link Google Sheets tidak valid.')
    return m.group(1)


def last_filled_row(svc, sid=TARGET_ID):
    """Baris terakhir yang terisi di kolom A tab ALL."""
    _, total = grid_rows(svc, sid, TARGET_TAB)
    last = 1
    for start in range(2, total + 1, 50000):
        end = min(start + 49999, total)
        vals = svc.values().get(spreadsheetId=sid, range=f"'{TARGET_TAB}'!A{start}:A{end}").execute().get('values', [])
        if vals:
            last = start + len(vals) - 1
    return last


# ---------------------------------------------------------------- brand
def load_cache():
    return json.loads(CACHE_FILE.read_text(encoding='utf-8')) if CACHE_FILE.exists() else {}


def sheet_mapping(rows):
    m, conflict = {}, set()
    for r in rows:
        if not r['brand'] or r['brand'] == UNKNOWN:
            continue
        if r['pid'] in m and m[r['pid']] != [r['brand'], r['sub']]:
            conflict.add(r['pid'])
        m[r['pid']] = [r['brand'], r['sub']]
    if conflict:
        print(f'PERINGATAN: {len(conflict)} Product ID punya brand/sub berbeda di List Produk '
              '(dipakai yang paling bawah).')
    return m


def read_filled(path):
    ws = openpyxl.load_workbook(path, data_only=True).worksheets[0]
    out = {}
    for r in list(ws.iter_rows(values_only=True))[1:]:
        pid, name, brand, sub = (cell_str(c) for c in (list(r) + [None] * 4)[:4])
        if pid:
            out[pid] = (name, brand, sub)
    return out


def write_template(path, missing, names, count, sheet_rows):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Isi'
    ws.append(['Product ID', 'Product Name', 'Brand', 'Sub Brand', 'Jumlah Baris'])
    for p in missing:
        ws.append([p, names[p], None, None, count[p]])
        ws.cell(ws.max_row, 1).number_format = '@'
    pairs = sorted({(r['brand'], r['sub']) for r in sheet_rows if r['brand'] and r['brand'] != UNKNOWN})
    ref = wb.create_sheet('Daftar')
    ref.append(['Brand', 'Sub Brand (semua)', 'Brand-Sub yang sudah ada'])
    brands = sorted({b for b, _ in pairs})
    subs = sorted({s for _, s in pairs if s})
    for i in range(max(len(brands), len(subs), len(pairs))):
        ref.append([brands[i] if i < len(brands) else None, subs[i] if i < len(subs) else None,
                    f'{pairs[i][0]} | {pairs[i][1]}' if i < len(pairs) else None])
    for col, lo, hi in (('C', 'A', len(brands)), ('D', 'B', len(subs))):
        dv = DataValidation(type='list', formula1=f'=Daftar!${lo}$2:${lo}${hi + 1}', allow_blank=True,
                            showErrorMessage=False)  # boleh ketik brand/sub baru
        ws.add_data_validation(dv)
        dv.add(f'{col}2:{col}{len(missing) + 1}')
    for col, w in zip('ABCDE', (24, 70, 16, 34, 13)):
        ws.column_dimensions[col].width = w
    wb.save(path)


def resolve(svc, pids, names, count, mapping_file):
    """-> (peta pid->[brand,sub], id baru yang harus ditambah ke List Produk)."""
    mapping = load_cache()
    missing = [p for p in pids if p not in mapping]
    sheet_rows, sheet_pids = [], set()
    if missing:
        print(f'{len(missing)} Product ID belum ada di cache, baca ulang "{BRAND_TAB}"...')
        sheet_rows, _ = read_brand_sheet(svc)
        sheet_pids = {r['pid'] for r in sheet_rows}
        mapping.update(sheet_mapping(sheet_rows))
        CACHE_FILE.write_text(json.dumps(mapping, ensure_ascii=False), encoding='utf-8')
        missing = [p for p in pids if p not in mapping]
    new_ids = []
    if missing and mapping_file.exists():
        filled = read_filled(mapping_file)
        for p in list(missing):
            if p in filled and filled[p][1] and filled[p][2]:
                name, brand, sub = filled[p]
                mapping[p] = [brand, sub]
                if p not in sheet_pids:
                    new_ids.append((p, name or names[p], brand, sub))
                missing.remove(p)
    if missing:
        if mapping_file.exists():
            print(f'\nMASIH ADA {len(missing)} Product ID yang Brand/Sub Brand-nya belum diisi di {mapping_file.name}:')
            for p in missing:
                print(f'  {p} | {names[p][:70]} ({count[p]} baris)')
            print('Isi dulu (atau hapus/rename file itu kalau isinya dari file raw lain), lalu jalankan lagi.')
        else:
            write_template(mapping_file, missing, names, count, sheet_rows)
            print(f'\n{len(missing)} Product ID belum punya brand. File isian dibuat:\n  {mapping_file}\n'
                  'Isi kolom Brand + Sub Brand, simpan, lalu jalankan perintah yang sama lagi.')
        sys.exit(2)
    return mapping, new_ids, {r['pid']: r['row'] for r in sheet_rows if r['brand'] == UNKNOWN}


# ---------------------------------------------------------------- rencana & tulis
def plan_rows(svc, rows, mapping, limit=None, sid=TARGET_ID):
    """Semua baris ditambahkan ke bawah ALL -> (indeks baris, baris pertama, baris terakhir sekarang,
    nilai A..Z yang akan ditulis). Dihitung ulang saat Apply supaya posisi selalu terbaru."""
    last_row = last_filled_row(svc, sid)
    new_idx = list(range(len(rows)))[:limit]
    first = last_row + 1
    out = [build_row(rows[i], first + j, *mapping[rows[i]['Product ID']]) for j, i in enumerate(new_idx)]
    return new_idx, first, last_row, out


def write_new_ids(svc, new_ids):
    """new_ids = [(pid, nama, brand, sub)] -> ditambah ke bawah List Produk (yang sudah ada dilewati)."""
    cur_rows, brand_last = read_brand_sheet(svc)
    have = {r['pid'] for r in cur_rows}
    new_ids = [x for x in new_ids if x[0] not in have]
    if not new_ids:
        return 0
    n0 = max((int(r['no']) for r in cur_rows if r['no'].isdigit()), default=0)
    ensure_rows(svc, BRAND_ID, BRAND_TAB, brand_last + len(new_ids))
    svc.values().update(
        spreadsheetId=BRAND_ID, range=f"'{BRAND_TAB}'!A{brand_last + 1}:E{brand_last + len(new_ids)}",
        valueInputOption='USER_ENTERED',
        body={'values': [[n0 + i + 1, txt(p), txt(name), brand, sub]
                         for i, (p, name, brand, sub) in enumerate(new_ids)]}).execute()
    print(f'{len(new_ids)} ID baru ditambah ke {BRAND_TAB}.')
    return len(new_ids)


def write_rows(svc, out, first, sid=TARGET_ID):
    if not out:
        return 0
    ensure_rows(svc, sid, TARGET_TAB, first + len(out) - 1)
    for s in range(0, len(out), 2000):
        chunk = out[s:s + 2000]
        svc.values().update(
            spreadsheetId=sid, range=f"'{TARGET_TAB}'!A{first + s}:Z{first + s + len(chunk) - 1}",
            valueInputOption='USER_ENTERED', body={'values': chunk}).execute()
        print(f'Tertulis {s + len(chunk)}/{len(out)} baris.')
    return len(out)


# ---------------------------------------------------------------- main
def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split('ALUR')[0].strip())
    ap.add_argument('file')
    ap.add_argument('--apply', action='store_true', help='benar-benar menulis ke Google Sheets')
    ap.add_argument('--limit', type=int, help='cuma tulis N baris baru pertama (buat cek tampilan)')
    ap.add_argument('--mapping', help='file isian Brand/Sub Brand (default Isi_Brand_<nama file>.xlsx)')
    a = ap.parse_args(argv)

    src = Path(a.file)
    rows = read_raw(src)
    print(f'{src.name}: {len(rows)} baris, tanpa filter.')
    names, count = {}, collections.Counter()
    for r in rows:
        names.setdefault(r['Product ID'], r['Product Name'])
        count[r['Product ID']] += 1
    mapping_file = Path(a.mapping) if a.mapping else src.with_name(f'Isi_Brand_{src.stem}.xlsx')

    svc = get_service(a.apply)
    mapping, new_ids, unknown_in_sheet = resolve(svc, list(names), names, count, mapping_file)
    for p, row in unknown_in_sheet.items():
        if p in names:
            print(f'CATATAN: {p} ada di List Produk baris {row} dengan brand "{UNKNOWN}"; '
                  'isian lo dipakai untuk import ini, tapi baris itu tidak diubah script (ubah manual).')

    rows, skipped = drop_skipped(rows, mapping)
    if skipped:
        print(f'Dilewati (Sub Brand {sorted(SKIP_SUBS)}): {len(skipped)} baris dari '
              f'{len({r["Product ID"] for r in skipped})} produk, tidak dimasukkan ke ALL.')
    print('Baca tab ALL untuk cari baris terakhir...')
    new_idx, first, last_row, out = plan_rows(svc, rows, mapping, a.limit)

    by_brand = collections.Counter(mapping[rows[i]['Product ID']][0] for i in new_idx)
    print(f'\nBaris baru: {len(new_idx)} -> baris {first}-{first + len(out) - 1} di ALL (baris terakhir sekarang {last_row})')
    print('Per brand:', dict(by_brand))
    print(f'ID baru untuk List Produk: {len(new_ids)}')
    for p, name, brand, sub in new_ids:
        print(f'  {p} | {name[:60]} | {brand} | {sub}')
    if not a.apply:
        print('\nDRY-RUN: tidak ada yang ditulis. Tambahkan --apply untuk menulis.')
        return
    if not out and not new_ids:
        print('\nTidak ada yang perlu ditulis.')
        return

    write_new_ids(svc, new_ids)
    write_rows(svc, out, first)
    CACHE_FILE.write_text(json.dumps(mapping, ensure_ascii=False), encoding='utf-8')
    print('\nSELESAI.')


if __name__ == '__main__':
    main()
